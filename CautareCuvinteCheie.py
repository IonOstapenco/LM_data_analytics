import os
import re
import PyPDF2
from docx import Document
from openpyxl import Workbook

# CONFIGURARE
folder_path = r"C:\Users\crme240\OneDrive - ION\Desktop\VMware"  # calea mapei cu PDF/DOCX

# cuvinte cheie
keywords = [
     "backup", "DR", "recover",
     "Disaster Recovery", "disaster recovery", "back-ups", "back"
]
output_file = os.path.join(folder_path, "rezultate_cautare.xlsx")

# Creăm un Excel nou
wb = Workbook()
ws = wb.active
ws.title = "Rezultate"
ws.append(["Fișier", "Pagina/Paragraf", "Cuvânt găsit", "Context"])


def search_in_pdf(file_path):
    results = []
    try:
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)

            # Dacă fișierul e criptat
            if reader.is_encrypted:
                try:
                    reader.decrypt("")  # încercăm fără parolă
                except Exception:
                    print(f" Fișier criptat, sărit: {os.path.basename(file_path)}")
                    results.append((os.path.basename(file_path), "N/A", "Fișier criptat", "Nu s-a putut citi conținutul"))
                    return results

            # Citim pagină cu pagină
            for page_num, page in enumerate(reader.pages, start=1):
                text = page.extract_text()
                if text:
                    for kw in keywords:
                        if re.search(rf"\b{kw}\b", text, re.IGNORECASE):
                            start = text.lower().find(kw.lower())
                            context = text[start:start+80].replace("\n", " ")
                            results.append((os.path.basename(file_path), f"Pagina {page_num}", kw, context))
    except Exception as e:
        print(f" Eroare la fișierul {os.path.basename(file_path)}: {e}")
        results.append((os.path.basename(file_path), "N/A", "Eroare", str(e)))
    return results


def search_in_docx(file_path):
    results = []
    try:
        doc = Document(file_path)
        for para_num, para in enumerate(doc.paragraphs, start=1):
            text = para.text
            for kw in keywords:
                if re.search(rf"\b{kw}\b", text, re.IGNORECASE):
                    start = text.lower().find(kw.lower())
                    context = text[start:start+80]
                    results.append((os.path.basename(file_path), f"Paragraf {para_num}", kw, context))
    except Exception as e:
        print(f"⚠️ Eroare la fișierul {os.path.basename(file_path)}: {e}")
        results.append((os.path.basename(file_path), "N/A", "Eroare", str(e)))
    return results


# Iterăm prin toate fișierele din mapă
for file in os.listdir(folder_path):
    full_path = os.path.join(folder_path, file)
    if file.lower().endswith(".pdf"):
        for r in search_in_pdf(full_path):
            ws.append(r)
    elif file.lower().endswith(".docx"):
        for r in search_in_docx(full_path):
            ws.append(r)

# Salvăm rezultatele în Excel
wb.save(output_file)
print(f"✅ Rezultatele au fost salvate în {output_file}")
