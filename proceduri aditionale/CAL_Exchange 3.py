import os
import csv
import json
import pandas as pd
from openpyxl import Workbook
from pathlib import Path

# ---------------------
# lettura del file parametri.json
with open("Parametri.json", encoding="utf-8") as f:
    config = json.load(f)

SOURCE_DIR = config["Source_dir"]
REPORT_DIR = config["Report_dir"]
OUTPUT_PATH = Path(SOURCE_DIR) / REPORT_DIR / config["Output_dir"]

#AD_SUBFOLDER = r"Active Directory Results - SERVIZICED - 2.12.2025-151753\ServiziCed.local"  # al momento hardcoded qui, non so ancora come inserirlo in parametri.json

#aggiunte in Parametri,json
AD_BASE_PATTERN = config["AD_Base_Pattern"]
AD_DOMAIN_SUBFOLDER = config["AD_Domain_Subfolder"]

# ---------------------
# cerchiamo automaticamento la più recente cartella AD

report_root = Path(SOURCE_DIR) / REPORT_DIR

ad_folders = [
    p for p in report_root.iterdir()
    if p.is_dir() and p.name.startswith(AD_BASE_PATTERN)
]

#se non lo e, visualiziamo il mesaggio
if not ad_folders:
    raise FileNotFoundError("! Atenzione! Non c'E nessuna cartella con quest nome. (Nu este niciun folder) Active Directory Results")

latest_ad_folder = max(ad_folders, key=lambda p: p.stat().st_mtime)

root_folder = latest_ad_folder / AD_DOMAIN_SUBFOLDER

print(f"[INFO] AD cartella scelta automaticamente/  AD folder ales automat: {root_folder}")


print(f"[DEBUG] Cautam in folder: {root_folder}")
print(f"[DEBUG] Exista folderul? {root_folder.exists()}")

# Listam fisierele gasite
all_csv_files = []
for root, dirs, files in os.walk(root_folder):
    for file in files:
        if file.endswith(".csv"):
            full_path = os.path.join(root, file)
            all_csv_files.append(full_path)
            print(f"[DEBUG] CSV gasit: {full_path}")

print(f"[DEBUG] Total fisiere CSV gasite: {len(all_csv_files)}")
if all_csv_files:
    print(f"[DEBUG] Primul CSV: {all_csv_files[0]}")


    
# ---------------------
# Output files
# Ho cambiato il nome in modo che la prossima volta inizi a elaborare i file che iniziano con AD

# output_file = os.path.join(SOURCE_DIR, REPORT_DIR, "usersAndGroupsResultFinale.csv")
#output_file_xlsx = os.path.join(SOURCE_DIR, REPORT_DIR, "Cal Exchenge.xlsx")
output_file_xlsx = OUTPUT_PATH / "Cal Exchenge.xlsx"

input_cn_file = os.path.join(
    SOURCE_DIR,
    REPORT_DIR,
    "output-1",
    config["AD_Users_Results_attivi"]
)

keywords_exclude = ["CR", "AX", "FI", "FJ", "test", "Administrator", "Cedacri_Admin", "harvest1", "PSMConnectCED", "R-xx1", "R-XZNETW"]


# normaliziamo keywords da excludere 
keywords_exclude_norm = [k.lower() for k in keywords_exclude]


# ---------------------
#legere files AD-usersAndGroupsResult.csv
# ---------------------
# ---------------------
# legere files AD-usersAndGroupsResult.csv
users = set()  # (UserName, SamAccountName)

for root, dirs, files in os.walk(root_folder):
    for file in files:
        if file == "AD-usersAndGroupsResult.csv":
            file_path = os.path.join(root, file)
            print(f"[DEBUG] Am gasit fisierul: {file_path}")
            
            try:
                with open(file_path, newline="", encoding="utf-16") as csvfile:   # revenim la utf-16 momentan
                    reader = csv.DictReader(csvfile, delimiter=";")
                    
                    print(f"[DEBUG] Coloane gasite: {reader.fieldnames}")
                    
                    if "UserName" not in reader.fieldnames or "SamAccountName" not in reader.fieldnames:
                        print(f"[ATTENZIONE] Coloana lipseste in {file_path}")
                        print(f"[ATTENZIONE] Coloane disponibile: {reader.fieldnames}")
                        continue

                    row_count = 0
                    added_count = 0
                    
                    for row in reader:
                        row_count += 1
                        username = str(row.get("UserName", "")).strip()
                        sam = str(row.get("SamAccountName", "")).strip()

                        if sam and any(sam.upper().startswith(p) for p in keywords_exclude):
                            continue

                        if username or sam:
                            users.add((username, sam))
                            added_count += 1
                        
                        if row_count <= 8:   # primele 8 randuri
                            print(f"[DEBUG] Rand {row_count}: UserName='{username}' | Sam='{sam}'")

                    print(f"[DEBUG] Total randuri citite: {row_count} | Adaugate in set: {added_count} | Total users: {len(users)}")

            except Exception as e:
                print(f"[ERROR] La citirea {file_path} -> {type(e).__name__}: {e}")

# ---------------------
#scrivere CSV
# scriem CSV
 #with open(output_file, "w", newline="", encoding="utf-8") as csvfile:
    #writer = csv.writer(csvfile)
    #writer.writerow(["UserName", "SamAccountName"])
    #for username, sam in sorted(users):
        #writer.writerow([username, sam])

#print(f"Fisier CSV generat: {output_file}")
#print(f"Total UserName unicale: {len(users)}")

# ---------------------
#scriviamo Excel

wb = Workbook()

#main Sheet -- il principale sheet
ws_main = wb.active
ws_main.title = "AD-usersAndGroupsResultFinale"
ws_main.append(["UserName", "SamAccountName"])

for username, sam in sorted(users):
    ws_main.append([username, sam])

# autosize
for col in ws_main.columns:
    max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
    ws_main.column_dimensions[col[0].column_letter].width = max_length + 2

# ---------------------
# sheet Results
ws_results = wb.create_sheet(title="ADattivi & ADuser")
ws_results.append(["CN_00", "UserName", "InnerJoin"])

# ---------------------
#legiamo CN_00
cn_list = []
try:
    cn_df = pd.read_csv(
        input_cn_file,
        sep="|",
        encoding="utf-8",
        skiprows=1,
        dtype=str
    ).fillna("")

    if "CN_00" not in cn_df.columns:
        print("[ERROR] Coloana CN_00 lipseste")
        cn_df = pd.DataFrame(columns=["CN_00"])
    else:
        cn_list = [cn.strip() for cn in cn_df["CN_00"] if cn.strip()]
        cn_df = pd.DataFrame({"CN_00": cn_list})

except Exception as e:
    print(f"[ERROR] Citire CN_00 -> {e}")
    cn_df = pd.DataFrame(columns=["CN_00"])

# ---------------------
# Users dataframe
# ---------------------
# Users dataframe
username_list = []
for u, s in sorted(users):
    u_clean = str(u).strip() if u is not None else ""
    if u_clean:
        username_list.append(u_clean)

users_df = pd.DataFrame({"UserName": username_list})

# ==================== DEBUG ADAUGAT AICI ====================
print("\n=== DEBUG users_df ===")
print("Tipul coloanei UserName:", users_df["UserName"].dtype)
print("Numar valori NaN:", users_df["UserName"].isna().sum())
print("Numar valori goale/string goale:", (users_df["UserName"].astype(str).str.strip() == "").sum())
print("\nPrimele 15 valori:")
print(users_df.head(15))
print("====================================\n")
# ============================================================

# Forțează string
users_df["UserName"] = users_df["UserName"].astype("string")   # sau .astype(str)
users_df["UserName_norm"] = users_df["UserName"].str.lower().str.strip()

# La fel pentru cn_df
cn_df["CN_00"] = cn_df["CN_00"].astype("string")
cn_df["CN_00_norm"] = cn_df["CN_00"].str.lower().str.strip()

# applichiamo esclusioni su UserName_norm
users_df = users_df[
    ~users_df["UserName_norm"].apply(
        lambda x: any(k in x for k in keywords_exclude_norm)
    )
]


# *-*-*-*-*-*
# INNER JOIN (case-insensitive)

merged_df = pd.merge(
    cn_df,
    users_df,
    left_on="CN_00_norm",
    right_on="UserName_norm",
    how="inner"
)


# creiamo una lista Users con i valori derivati dall’Inner Join






'''
#cream inner join ordonat
# doar INNER JOIN + sortare alfabetica
merged_df_sorted = (
    merged_df
    .loc[:, ["CN_00", "UserName"]]
    .dropna()
    .drop_duplicates()
    .sort_values(by="UserName", key=lambda col: col.str.lower())
)
'''



print(f"Numar recorduri INNER JOIN: {len(merged_df)}")
print(f"Numero di record INNER JOIN: {len(merged_df)}")

# ---------------------
# Results
max_len = max(len(cn_list), len(username_list))

for i in range(max_len):
    cn_val = cn_list[i] if i < len(cn_list) else ""
    username_val = username_list[i] if i < len(username_list) else ""

    inner_val = (
        cn_val
        if cn_val.lower() in merged_df["CN_00_norm"].values
        else ""
    )


    ws_results.append([cn_val, username_val, inner_val])



    

# autosize Results
for col in ws_results.columns:
    max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
    ws_results.column_dimensions[col[0].column_letter].width = max_length + 2


# ---------------------
# sheet user Cal_Exchenge - DOAR INNER JOIN pe UserName
ws_results2 = wb.create_sheet(title="user Cal_Exchenge")
ws_results2.append(["USERS"])

# creazione della lista dei UserName risultante dall’INNER JOIN, ordinata alfabeticamente

users_innerjoin_sorted = merged_df["UserName"].dropna().drop_duplicates().sort_values(key=lambda col: col.str.lower())

for user in users_innerjoin_sorted:
    ws_results2.append([user])

# autosize 
for col in ws_results2.columns:
    max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
    ws_results2.column_dimensions[col[0].column_letter].width = max_length + 2




''' 
# ---------------------
# sheet Results2 - DOAR INNER JOIN
ws_results2 = wb.create_sheet(title="Results2")
ws_results2.append(["CN_00", "USERS"])

for _, row in merged_df_sorted.iterrows():
    ws_results2.append([row["CN_00"], row["UserName"]])

# autosize Results2
for col in ws_results2.columns:
    max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
    ws_results2.column_dimensions[col[0].column_letter].width = max_length + 2   


'''    

# ---------------------
#salvataggio Excel
# salvare Excel
wb.save(output_file_xlsx)
print(f"File Excel generato: {output_file_xlsx}")
