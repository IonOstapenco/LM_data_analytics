import os
import csv
import json
from openpyxl import Workbook

# functie de transformare csv in exce;
def flexera_csv_to_xlsx(fisier):
    with open(fisier, "r", encoding="utf-8") as f:
        cfg = json.load(f)

    csv_file = os.path.join(
        cfg["Source_dir"],
        cfg["Report_dir"],
        cfg["FLEXERA_SOFTWARE_CSV"]
    )

    # fisier exlsx
    xlsx_fisier = os.path.splitext(csv_file)[0] + cfg["XLS_end"]


    #VARIANTA VECHE
    #  care genereaza 2 sheeturi, 1 cu default sheet, care este gol, si unul cu info, "Flexera"
    #wb = Workbook(write_only=False) # --> varianta veche
    # se numeste software (1) edobicei, cand transformam manual
    #ws = wb.create_sheet("Flexera")

    wb = Workbook()
    ws = wb.active
    ws.title = "Flexera"

    with open(csv_file, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f, delimiter=",")

        for nr, row in enumerate(reader, start=1):
            ws.append(row)

            #DEBUG
            if nr % 10000 == 0:
                print(f"{nr}:, randuri convertite ,...")
    wb.save(xlsx_fisier)

    print("conversie terminata,,,")
    print(xlsx_fisier)

if __name__ == "__main__":
    flexera_csv_to_xlsx("parametri.json")                




